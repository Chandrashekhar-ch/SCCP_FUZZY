"""
Generates a standalone introduction draft that narrates the full project
lifecycle, two paragraphs per step:
  1. Problem & Motivation
  2. Sensor & Hardware Selection
  3. Preprocessing
  4. Fuzzy Inference System Design
  5. Implementation (Dashboard & Deployment)
  6. Validation & Results
  7. Conclusion & Outlook

Run from project root: python docs/generate_intro_draft.py
Output: outputs/Project_Introduction_Draft.docx
"""
import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Calibri"


def set_font(run, size=11.5, bold=False, italic=False, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def add_heading(doc, text, level=1, size=None, center=False):
    h = doc.add_heading(level=level)
    r = h.add_run(text)
    set_font(r, size=size or {1: 16, 2: 13.5}.get(level, 12), bold=True)
    if center:
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    return h


def add_para(doc, text, size=11.5, bold=False, italic=False, center=False, space_after=10):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, italic=italic)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_step(doc, num, title, para1, para2):
    add_heading(doc, f"Step {num}: {title}", level=2)
    add_para(doc, para1)
    add_para(doc, para2)


def create_intro_draft():
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    title = doc.add_heading(level=0)
    tr = title.add_run("Project Introduction: A Fuzzy Inference System for "
                        "Real-Time Water Potability Classification")
    set_font(tr, size=18, bold=True)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, "Project Flow Overview — SC_CProject", size=12, italic=True, center=True, space_after=20)

    add_para(doc, "This document introduces the project by walking through its complete lifecycle, from "
                   "the original motivation through to the validated outcome. Each stage of the project is "
                   "presented as a self-contained step, described in two paragraphs: the first sets out "
                   "what the step does and why it exists in the pipeline, and the second describes how it "
                   "was actually carried out in this project, including the concrete design choices made.",
             italic=True, space_after=20)

    # Step 1
    add_step(doc, 1, "Problem & Motivation",
        "Access to safe drinking water is a basic public-health requirement, yet verifying that a given "
        "water source is safe still typically depends on collecting a sample and sending it to a "
        "laboratory. That process is accurate, but it is slow, costly, and cannot realistically be "
        "repeated often enough or at enough locations to give anyone a real-time picture of water "
        "quality. Low-cost electronic sensors exist for the key parameters that determine potability, "
        "but the simplest way of using them — comparing each reading to a fixed numeric threshold — "
        "creates an unrealistic “cliff-edge” decision: a reading just inside a safe limit is treated as "
        "perfectly fine, while a reading a hair's breadth outside it is treated as a total failure, even "
        "though the physical difference between the two is negligible.",
        "This project starts from that gap: the need for a monitoring approach that is both affordable "
        "enough to run continuously and expressive enough to avoid the cliff-edge problem. The chosen "
        "direction was fuzzy logic, specifically a Mamdani-type Fuzzy Inference System (FIS), because it "
        "represents each parameter using graded, overlapping categories (for example, a pH reading can be "
        "partly “neutral” and partly “slightly acidic” at the same time) rather than a single hard cutoff. "
        "Just as importantly, the project deliberately avoided a black-box machine-learning classifier as "
        "the core decision-maker, because every fuzzy rule used is written in plain IF–THEN language and "
        "traced back to a specific WHO or BIS drinking-water guideline, so every classification the system "
        "produces can be explained and audited rather than simply trusted.")

    # Step 2
    add_step(doc, 2, "Sensor & Hardware Selection",
        "Before any inference logic could be designed, the project had to settle on which physical "
        "parameters to measure and which low-cost sensors could measure them reliably. Water potability "
        "is not determined by a single number; it depends on a combination of chemical and physical "
        "properties, so the sensor set had to cover enough independent dimensions of water quality to "
        "support meaningful, multi-parameter rules, while staying within a hobbyist/prototype budget and "
        "remaining compatible with a single microcontroller.",
        "The project settled on four potability-relevant sensors — pH, turbidity, total dissolved solids "
        "(TDS), and dissolved oxygen (DO) — plus a DS18B20 temperature probe used purely as a calibration "
        "reference rather than a potability parameter in its own right. All four sensors are from the "
        "DFRobot Gravity series so that they share a common analog/digital interface, and together with "
        "the temperature probe they bring the total sensor cost to roughly ₹5,350. An ESP32 "
        "microcontroller was chosen as the acquisition front end for its built-in 12-bit ADC and Wi-Fi "
        "radio, letting it sample all five channels and stream the readings onward at 115,200 baud "
        "without any additional networking hardware.")

    # Step 3
    add_step(doc, 3, "Preprocessing",
        "Raw sensor readings are never clean enough to feed directly into a decision system: pH and "
        "dissolved-oxygen probes drift with water temperature in well-understood but non-trivial ways, "
        "and any analog sensor is subject to momentary noise spikes and occasional bad readings that, if "
        "left uncorrected, would show up as false alarms or missed detections in the final "
        "classification. A dedicated preprocessing stage was therefore placed between the raw sensor "
        "stream and the inference engine to make sure the values the fuzzy system sees are physically "
        "meaningful.",
        "In this project, preprocessing performs three corrections in sequence. First, temperature "
        "compensation is applied: the pH reading is adjusted using the relation ΔpH = 0.003 × (T − 25°C) "
        "to correct for Nernstian drift, and the DO reading is rescaled using the Benson–Krause "
        "temperature-solubility relation, DO_corrected = DO_raw × (DO_sat(25°C) / DO_sat(T)). Second, "
        "statistical outlier rejection (Z-score based) discards individual readings that are implausible "
        "given recent history. Third, a moving-average filter over the last 10 samples smooths out "
        "residual sensor noise, producing the clean, stable crisp values that are handed to the fuzzy "
        "inference engine.")

    # Step 4
    add_step(doc, 4, "Fuzzy Inference System Design",
        "This step is the analytical core of the project: turning four cleaned sensor readings into a "
        "single, trustworthy verdict on water potability. Rather than compare each reading to a fixed "
        "limit, the system needed a way to reason the way a human expert would — treating a borderline "
        "reading as borderline, weighing several moderately-off parameters together, and still being able "
        "to explain exactly why a given sample was judged safe, marginal, or unsafe.",
        "The project implements this as a four-stage Mamdani pipeline. Each input is first fuzzified into "
        "degrees of membership across a small set of overlapping linguistic terms (for example pH is "
        "described using Acidic, Slightly Acidic, Neutral, Slightly Alkaline and Alkaline). These "
        "memberships are then run through a rule base of 45 IF–THEN rules — every one of them grounded in "
        "a specific WHO (2022) or BIS IS 10500:2012 limit — combined with the MIN operator and aggregated "
        "across rules with the MAX operator. Finally, the aggregated fuzzy result is defuzzified using the "
        "centroid-of-area method to produce a single crisp Water Potability Index (WPI) on a 0–10 scale, "
        "which is then classified into Non-Potable, Marginal, or Potable.")

    # Step 5
    add_step(doc, 5, "Implementation (Dashboard & Deployment)",
        "A classification number is only useful if someone can see it, understand it, and trust it in the "
        "moment. The project therefore needed an implementation layer that turns the FIS engine into "
        "something an operator can actually watch and interact with in real time, rather than a script "
        "that only prints a final number.",
        "This was built as a Flask + Socket.IO web application that streams live updates to a browser "
        "dashboard. The dashboard shows an animated WPI arc gauge, live status cards for each of the four "
        "sensors, a rolling 60-second WPI trend chart, live visualizations of the current fuzzy membership "
        "degrees, and a panel showing exactly which of the 45 rules are firing and how strongly — making "
        "the “why” behind every classification directly visible. It also includes a manual, slider-driven "
        "inference mode and a selector across six demonstration scenarios (normal, ideal potable, acidic, "
        "turbid, hard water, and worst-case contaminated), so the system can be explored and demonstrated "
        "even without live hardware attached.")

    # Step 6
    add_step(doc, 6, "Validation & Results",
        "A system built on expert-written rules still needs to be checked against data, both to confirm "
        "it behaves sensibly and to see how it compares with more conventional approaches. The project's "
        "validation step was designed to answer two separate questions: how does the FIS perform against "
        "the very standard it was built from, and how does it perform on messy, real-world data it has "
        "never seen, compared against trained machine-learning classifiers.",
        "For the first question, a 5,000-sample synthetic dataset was generated and labelled by a "
        "WHO/BIS hard-threshold oracle; here the FIS reached 75.3% accuracy (75.2% macro-F1), below "
        "Random Forest (99.5%), SVM (95.7%) and KNN (92.3%) — an expected result, since those models were "
        "trained directly against the same oracle used to create the labels. For the second, more telling "
        "question, all four methods were evaluated on the independent, real-world Kaggle water-potability "
        "dataset: the FIS scored 60.43% accuracy, ahead of Random Forest (59.61%) and KNN (54.40%) and "
        "close behind SVM (61.22%) — despite having been designed with zero exposure to that dataset's "
        "labels, and while remaining the only one of the four methods whose every decision can be traced "
        "back to a specific, named rule.")

    # Step 7
    add_step(doc, 7, "Conclusion & Outlook",
        "Taken together, the seven steps above form a complete path from an unmet real-world need to a "
        "working, evaluated system: a clearly identified gap in low-cost water monitoring, a concrete "
        "hardware choice to address it, a preprocessing stage to make the hardware trustworthy, a "
        "standards-grounded fuzzy reasoning core to make the decisions interpretable, a live dashboard to "
        "make those decisions visible, and a validation study to check the whole system's performance "
        "honestly, including on data it was never designed around.",
        "The outcome is a system that gives up a modest amount of raw accuracy, compared to black-box "
        "machine-learning models, in exchange for full traceability of every classification back to a "
        "published drinking-water standard — a trade the project considers justified for a safety-related "
        "decision that affects public trust. Documented next steps include field-testing the sensor "
        "hardware on real water sources, exploring hybrid tuning of the fuzzy rules against expert-"
        "labelled field data without losing interpretability, and moving the inference engine onto the "
        "ESP32 itself so the system can run without a separate host computer.")

    os.makedirs("outputs", exist_ok=True)
    out_path = "outputs/Project_Introduction_Draft.docx"
    try:
        doc.save(out_path)
    except PermissionError:
        out_path = "outputs/Project_Introduction_Draft_new.docx"
        doc.save(out_path)
    print(f"[*] Introduction draft saved to {out_path}")


if __name__ == "__main__":
    create_intro_draft()
