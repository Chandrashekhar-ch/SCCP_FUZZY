"""
Generates the project presentation deck for:
Fuzzy Inference System for Real-Time Water Potability Classification

Run from project root:  python docs/generate_ppt.py
Output: outputs/Water_Potability_FIS_Presentation.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

# ── Theme ──────────────────────────────────────────────────────────────
NAVY = RGBColor(0x0B, 0x2A, 0x3D)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
TEAL_LT = RGBColor(0xE3, 0xF3, 0xF3)
ACCENT = RGBColor(0x22, 0xC5, 0x5E)      # potable green
WARN = RGBColor(0xF5, 0x9E, 0x0B)        # marginal amber
DANGER = RGBColor(0xEF, 0x44, 0x44)      # non-potable red
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x55, 0x62, 0x6B)
LIGHT_BG = RGBColor(0xF7, 0xFA, 0xFA)

FIG_DIR = "outputs/figures"
SW, SH = Inches(13.333), Inches(7.5)  # 16:9


def add_background(slide, color=WHITE):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    bg.shadow.inherit = False
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)
    return bg


def add_textbox(slide, left, top, width, height, text, size=18, bold=False,
                 color=NAVY, align=PP_ALIGN.LEFT, font="Calibri", italic=False,
                 line_spacing=1.0):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return box


def add_bullets(slide, left, top, width, height, items, size=16, color=NAVY,
                 font="Calibri"):
    """items: list of (text, level, bold)"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        text, level, bold = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        p.space_after = Pt(10 if level == 0 else 6)
        run = p.add_run()
        bullet_char = "▸ " if level == 0 else "–  "
        run.text = f"{bullet_char}{text}"
        run.font.size = Pt(size - level * 2)
        run.font.bold = bold
        run.font.color.rgb = color if level == 0 else GREY
        run.font.name = font
    return box


def add_header(slide, title_text, subtitle=None, kicker=None):
    add_background(slide, WHITE)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), SH)
    bar.fill.solid(); bar.fill.fore_color.rgb = TEAL; bar.line.fill.background()
    bar.shadow.inherit = False

    if kicker:
        add_textbox(slide, Inches(0.6), Inches(0.35), Inches(8), Inches(0.4),
                    kicker.upper(), size=13, bold=True, color=TEAL, font="Calibri")
    add_textbox(slide, Inches(0.55), Inches(0.68 if kicker else 0.5), Inches(11.5), Inches(0.9),
                title_text, size=30, bold=True, color=NAVY, font="Calibri")
    if subtitle:
        add_textbox(slide, Inches(0.6), Inches(1.35), Inches(11.5), Inches(0.5),
                    subtitle, size=15, italic=True, color=GREY)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.28), Inches(2.2), Pt(3))
    rule.fill.solid(); rule.fill.fore_color.rgb = TEAL; rule.line.fill.background()
    rule.shadow.inherit = False
    return slide


def add_footer(slide, idx, total, section=""):
    add_textbox(slide, Inches(0.55), Inches(7.12), Inches(6), Inches(0.3),
                "Fuzzy Inference System · Water Potability Classification", size=9, color=GREY)
    add_textbox(slide, Inches(11.8), Inches(7.12), Inches(1.0), Inches(0.3),
                f"{idx}/{total}", size=9, color=GREY, align=PP_ALIGN.RIGHT)


def set_cell(cell, text, size=11, bold=False, color=NAVY, fill=None, align=PP_ALIGN.LEFT):
    cell.text = text
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Pt(6); cell.margin_right = Pt(6)
    cell.margin_top = Pt(3); cell.margin_bottom = Pt(3)
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = WHITE


def add_table(slide, left, top, width, height, rows_data, col_widths=None,
              header_fill=TEAL, header_color=WHITE, body_size=11, header_size=12,
              zebra=True):
    nrows = len(rows_data)
    ncols = len(rows_data[0])
    gtable = slide.shapes.add_table(nrows, ncols, left, top, width, height).table
    if col_widths:
        for i, w in enumerate(col_widths):
            gtable.columns[i].width = w
    for r, row in enumerate(rows_data):
        for c, val in enumerate(row):
            cell = gtable.cell(r, c)
            if r == 0:
                set_cell(cell, str(val), size=header_size, bold=True, color=header_color, fill=header_fill,
                         align=PP_ALIGN.CENTER)
            else:
                fill = LIGHT_BG if (zebra and r % 2 == 0) else WHITE
                set_cell(cell, str(val), size=body_size, color=NAVY, fill=fill)
    return gtable


def add_pill(slide, left, top, width, height, text, fill, text_color=WHITE, size=13):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shp.adjustments[0] = 0.5
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = text_color
    return shp


def add_pic_framed(slide, path, left, top, width=None, height=None, caption=None):
    if not os.path.exists(path):
        return
    pic = slide.shapes.add_picture(path, left, top, width=width, height=height)
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, pic.left, pic.top, pic.width, pic.height)
    frame.fill.background()
    frame.line.color.rgb = RGBColor(0xDD, 0xE3, 0xE6)
    frame.line.width = Pt(1)
    frame.shadow.inherit = False
    if caption:
        add_textbox(slide, pic.left, pic.top + pic.height + Inches(0.05), pic.width, Inches(0.3),
                    caption, size=11, italic=True, color=GREY, align=PP_ALIGN.CENTER)
    return pic


# ══════════════════════════════════════════════════════════════════════
def create_ppt():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    blank = prs.slide_layouts[6]
    TOTAL = 16

    # ── Slide 1: Title ──────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    add_background(s, NAVY)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(5.35), SW, Inches(2.15))
    band.fill.solid(); band.fill.fore_color.rgb = TEAL; band.line.fill.background()
    band.shadow.inherit = False
    add_textbox(s, Inches(0.9), Inches(1.65), Inches(11.5), Inches(0.5),
                "SC_CPROJECT · FUZZY LOGIC SYSTEMS", size=15, bold=True, color=RGBColor(0x7D, 0xD3, 0xD8))
    add_textbox(s, Inches(0.85), Inches(2.15), Inches(11.6), Inches(1.9),
                "Fuzzy Inference System for Real-Time\nWater Potability Classification", size=38, bold=True, color=WHITE,
                line_spacing=1.05)
    add_textbox(s, Inches(0.9), Inches(4.05), Inches(11), Inches(0.7),
                "From Low-Cost Multi-Sensor Data — a Mamdani FIS that classifies drinking-water\n"
                "safety in real time using four low-cost IoT sensors, no lab equipment required.",
                size=15, italic=True, color=RGBColor(0xCF, 0xE8, 0xE9), line_spacing=1.2)
    add_textbox(s, Inches(0.9), Inches(5.75), Inches(8), Inches(0.5),
                "Presented by: SC_CProject Team", size=16, bold=True, color=WHITE)
    add_textbox(s, Inches(0.9), Inches(6.2), Inches(8), Inches(0.5),
                "WHO Guidelines 2022  ·  BIS IS 10500:2012", size=13, color=RGBColor(0xE0, 0xF0, 0xF1))

    # ── Slide 2: Problem Statement ─────────────────────────────────
    s = prs.slides.add_slide(blank)
    add_header(s, "Problem Statement & Motivation", kicker="Introduction")
    add_bullets(s, Inches(0.65), Inches(1.7), Inches(6.4), Inches(4.8), [
        ("Conventional water testing is slow & costly", 0, True),
        ("Lab analysis requires trained personnel, days of turnaround, and is geographically limited.", 1, False),
        ("Hard-threshold classifiers create a “cliff-edge” effect", 0, True),
        ("A tiny change in a sensor reading can flip the decision abruptly at a boundary — unrealistic for real water.", 1, False),
        ("Black-box ML models are not auditable", 0, True),
        ("Neural networks give no traceable reasoning — unacceptable for safety-critical drinking-water decisions.", 1, False),
        ("Need: real-time, low-cost, interpretable classification", 0, True),
    ], size=16)
    # Right-side stat cards
    stats = [("4", "Low-cost IoT\nsensors", TEAL), ("45", "WHO/BIS grounded\nfuzzy rules", ACCENT), ("<₹8,000", "Total hardware\ncost", WARN)]
    x = Inches(7.5)
    for i, (num, label, col) in enumerate(stats):
        y = Inches(1.75) + i * Inches(1.55)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.1), Inches(1.3))
        card.adjustments[0] = 0.08
        card.fill.solid(); card.fill.fore_color.rgb = TEAL_LT; card.line.color.rgb = col; card.line.width = Pt(1.25)
        card.shadow.inherit = False
        add_textbox(s, x + Inches(0.25), y + Inches(0.12), Inches(1.4), Inches(1.0), num, size=34, bold=True, color=col)
        add_textbox(s, x + Inches(1.7), y + Inches(0.28), Inches(3.2), Inches(0.8), label, size=14, color=NAVY, line_spacing=1.1)
    add_footer(s, 2, TOTAL)

    # ── Slide 3: Why Fuzzy Logic (comparison) ──────────────────────
    s = prs.slides.add_slide(blank)
    add_header(s, "Why Fuzzy Logic?", kicker="Introduction")
    rows = [
        ["Approach", "Weakness / Strength"],
        ["Hard-threshold classifier", "Cliff-edge effect at boundaries — tiny change causes drastic class jump"],
        ["Neural network (black box)", "High accuracy but not interpretable — cannot be audited for safety-critical use"],
        ["Mamdani FIS (this project)", "Graded membership, human-readable rules, WHO/BIS grounded, real-time capable"],
    ]
    add_table(s, Inches(0.65), Inches(1.85), Inches(12.0), Inches(2.2), rows,
              col_widths=[Inches(3.6), Inches(8.4)], body_size=14, header_size=14)
    add_textbox(s, Inches(0.65), Inches(4.35), Inches(11.5), Inches(0.4),
                "Design Philosophy", size=16, bold=True, color=TEAL)
    add_bullets(s, Inches(0.65), Inches(4.8), Inches(11.7), Inches(2.0), [
        ("Every decision must be traceable to a specific rule grounded in a WHO/BIS standard.", 0, False),
        ("Outputs should be continuous (0–10 WPI) rather than a binary safe/unsafe label.", 0, False),
        ("The system must run on low-cost edge hardware in real time — no cloud inference needed.", 0, False),
    ], size=15)
    add_footer(s, 3, TOTAL)

    # ── Slide 4: Objectives ─────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    add_header(s, "Project Objectives", kicker="Introduction")
    objs = [
        "Design a Mamdani Fuzzy Inference System for water potability using pH, Turbidity, TDS & Dissolved Oxygen",
        "Ground every membership function and rule in WHO (2022) / BIS IS 10500:2012 standards",
        "Apply temperature-compensated preprocessing to correct sensor drift (pH, DO)",
        "Build a real-time web dashboard for live visualization of readings, membership, and rule activation",
        "Validate the FIS against ML classifiers (Random Forest, KNN, SVM) on synthetic & real-world datasets",
        "Demonstrate deployability on a ₹5,350 sensor stack + ESP32 microcontroller",
    ]
    items = [(o, 0, False) for o in objs]
    add_bullets(s, Inches(0.75), Inches(1.85), Inches(11.7), Inches(5.0), items, size=18)
    add_footer(s, 4, TOTAL)

    # ── Slide 5: System Architecture ────────────────────────────────
    s = prs.slides.add_slide(blank)
    add_header(s, "System Architecture", subtitle="Sensor → Preprocessing → Mamdani FIS → Real-Time Dashboard", kicker="Design")
    stages = [
        ("SENSOR LAYER", "pH · Turbidity · TDS · DO · Temperature\n(compensator)", TEAL),
        ("PREPROCESSING", "Temp. compensation → Z-score outlier\nrejection → Moving avg (N=10) → Clamp", NAVY),
        ("MAMDANI FIS ENGINE", "Fuzzification → 45-Rule Evaluation →\nMIN Implication → MAX Aggregation → Centroid Defuzz.", ACCENT),
        ("REAL-TIME DASHBOARD", "Flask + Socket.IO — Arc gauge, sensor\ncards, membership & rule visualisation", WARN),
    ]
    y = Inches(1.95)
    for i, (name, desc, col) in enumerate(stages):
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), y, Inches(11.9), Inches(1.05))
        box.adjustments[0] = 0.12
        box.fill.solid(); box.fill.fore_color.rgb = LIGHT_BG
        box.line.color.rgb = col; box.line.width = Pt(1.5)
        box.shadow.inherit = False
        add_textbox(s, Inches(0.95), y + Inches(0.12), Inches(3.3), Inches(0.8), name, size=16, bold=True, color=col)
        add_textbox(s, Inches(4.4), y + Inches(0.1), Inches(7.9), Inches(0.85), desc, size=13, color=GREY, line_spacing=1.15)
        if i < len(stages) - 1:
            arrow = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.1), y + Inches(1.03), Inches(0.35), Inches(0.22))
            arrow.fill.solid(); arrow.fill.fore_color.rgb = TEAL; arrow.line.fill.background()
            arrow.shadow.inherit = False
        y += Inches(1.28)
    add_footer(s, 5, TOTAL)

    # ── Slide 6: Sensors & Hardware ─────────────────────────────────
    s = prs.slides.add_slide(blank)
    add_header(s, "Sensor Selection & Hardware", kicker="Design")
    rows = [
        ["#", "Parameter", "Sensor", "Cost (INR)", "WHO/BIS Limit"],
        ["1", "pH", "DFRobot SEN0161", "~₹1,200", "6.5 – 8.5"],
        ["2", "Turbidity", "DFRobot SEN0189", "~₹800", "<1 NTU (WHO), <5 NTU (BIS)"],
        ["3", "TDS", "DFRobot SEN0244", "~₹700", "<500 mg/L (BIS desirable)"],
        ["4", "Dissolved Oxygen", "DFRobot SEN0237", "~₹2,500", ">6 mg/L (health benchmark)"],
        ["5", "Temperature", "DS18B20 (compensator)", "~₹150", "Calibration use only"],
    ]
    add_table(s, Inches(0.65), Inches(1.8), Inches(12.0), Inches(2.6), rows,
              col_widths=[Inches(0.6), Inches(2.6), Inches(3.3), Inches(2.2), Inches(3.3)],
              body_size=13, header_size=13)
    add_pill(s, Inches(0.65), Inches(4.65), Inches(3.6), Inches(0.55), "Total cost: ~₹5,350", ACCENT, size=15)
    add_textbox(s, Inches(0.65), Inches(5.45), Inches(11.7), Inches(0.4), "Microcontroller & Compensation", size=16, bold=True, color=TEAL)
    add_bullets(s, Inches(0.65), Inches(5.9), Inches(11.7), Inches(1.4), [
        ("ESP32 (Wi-Fi, 12-bit ADC, 240 MHz dual-core) streams JSON packets at 115,200 baud", 0, False),
        ("pH: ΔpH = 0.003 × (T − 25°C)   |   DO: corrected via Benson & Krause saturation model", 0, False),
    ], size=13)
    add_footer(s, 6, TOTAL)

    # ── Slide 7: FIS Design Overview ─────────────────────────────────
    s = prs.slides.add_slide(blank)
    add_header(s, "Fuzzy Inference System Design", subtitle="Mamdani-type FIS", kicker="FIS Design")
    pills = [("Implication", "MIN", TEAL), ("Aggregation", "MAX", NAVY), ("Defuzzification", "Centroid of Area", ACCENT)]
    x = Inches(0.65)
    for label, val, col in pills:
        add_textbox(s, x, Inches(1.85), Inches(3.9), Inches(0.35), label.upper(), size=12, bold=True, color=GREY)
        add_pill(s, x, Inches(2.2), Inches(3.9), Inches(0.65), val, col, size=17)
        x += Inches(4.1)
    rows = [
        ["Variable", "Range", "Linguistic Terms"],
        ["pH (input)", "0 – 14", "Acidic · Slightly Acidic · Neutral · Slightly Alkaline · Alkaline"],
        ["Turbidity (input)", "0 – 100 NTU", "Clear · Slightly Turbid · Turbid · Very Turbid"],
        ["TDS (input)", "0 – 1500 mg/L", "Pure · Acceptable · High · Very High"],
        ["DO (input)", "0 – 20 mg/L", "Very Low · Low · Acceptable · High"],
        ["WPI (output)", "0 – 10", "Non-Potable · Marginal · Potable"],
    ]
    add_table(s, Inches(0.65), Inches(3.15), Inches(12.0), Inches(3.3), rows,
              col_widths=[Inches(2.3), Inches(1.9), Inches(7.8)], body_size=13, header_size=13)
    add_footer(s, 7, TOTAL)

    # ── Slide 8: Input Membership Functions (figure) ────────────────
    s = prs.slides.add_slide(blank)
    add_header(s, "Fuzzy Sets: pH Membership Functions", kicker="FIS Design")
    add_pic_framed(s, f"{FIG_DIR}/mf_ph.png", Inches(1.7), Inches(1.75), width=Inches(9.9),
                   caption="pH universe of discourse [0,14] with 5 overlapping fuzzy sets; green band = WHO safe zone (6.5–8.5)")
    add_footer(s, 8, TOTAL)

    # ── Slide 9: Output Membership Function (figure) ────────────────
    s = prs.slides.add_slide(blank)
    add_header(s, "Output Variable: Water Potability Index (WPI)", kicker="FIS Design")
    add_pic_framed(s, f"{FIG_DIR}/mf_output_wpi.png", Inches(1.7), Inches(1.75), width=Inches(9.9),
                   caption="Three overlapping output sets on [0,10]: Non-Potable, Marginal, Potable")
    add_footer(s, 9, TOTAL)

    # ── Slide 10: Rule Base ──────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    add_header(s, "Rule Base — 45 WHO/BIS-Grounded Rules", kicker="FIS Design")
    rows = [
        ["Rule Group", "# Rules", "Consequent"],
        ["Fully potable (all parameters safe)", "8", "POTABLE"],
        ["Hard non-potability triggers (single-param failure)", "5", "NON-POTABLE"],
        ["Dual-parameter failures", "4", "NON-POTABLE"],
        ["Single-parameter boundary violation", "13", "MARGINAL"],
        ["Turbid + mixed parameters", "4", "MARGINAL / NON-POTABLE"],
        ["Extreme scenarios", "3", "NON-POTABLE"],
        ["Multi-parameter boundary cases", "8", "MARGINAL"],
    ]
    add_table(s, Inches(0.65), Inches(1.8), Inches(12.0), Inches(3.1), rows,
              col_widths=[Inches(7.3), Inches(1.7), Inches(3.0)], body_size=13, header_size=13)
    add_textbox(s, Inches(0.65), Inches(5.1), Inches(11.7), Inches(0.35), "Sample Rules", size=15, bold=True, color=TEAL)
    code_box = s.shapes.add_textbox(Inches(0.65), Inches(5.5), Inches(12.0), Inches(1.7))
    tf = code_box.text_frame
    tf.word_wrap = True
    lines = [
        "R1:  IF pH=NEUTRAL AND Turbidity=CLEAR AND TDS=ACCEPTABLE AND DO=HIGH  →  WPI=POTABLE  [w=1.0]",
        "R9:  IF pH=ACIDIC  →  WPI=NON_POTABLE  [w=1.0]   (critical trigger)",
        "R18: IF pH=NEUTRAL AND Turbidity=SLIGHTLY_TURBID AND TDS=ACCEPTABLE AND DO=ACCEPTABLE  →  WPI=MARGINAL  [w=0.80]",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run = p.add_run(); run.text = line
        run.font.name = "Consolas"; run.font.size = Pt(13); run.font.color.rgb = NAVY
    add_footer(s, 10, TOTAL)

    # ── Slide 11: Defuzzification ─────────────────────────────────────
    s = prs.slides.add_slide(blank)
    add_header(s, "Inference Pipeline & Defuzzification", kicker="FIS Design")
    add_bullets(s, Inches(0.65), Inches(1.85), Inches(6.0), Inches(4.5), [
        ("Stage 1 — Fuzzification", 0, True),
        ("Compute μ(x) for every linguistic term of each crisp input.", 1, False),
        ("Stage 2 — Rule Evaluation", 0, True),
        ("Strength = MIN(antecedent μ values) × rule weight, across 45 rules.", 1, False),
        ("Stage 3 — Implication + Aggregation", 0, True),
        ("Clip each output MF at its rule strength (MIN); combine all via MAX.", 1, False),
        ("Stage 4 — Defuzzification", 0, True),
        ("Centroid of Area yields a single crisp WPI ∈ [0, 10].", 1, False),
    ], size=15)
    formula_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(2.3), Inches(5.5), Inches(2.0))
    formula_box.adjustments[0] = 0.08
    formula_box.fill.solid(); formula_box.fill.fore_color.rgb = NAVY; formula_box.line.fill.background()
    formula_box.shadow.inherit = False
    add_textbox(s, Inches(7.35), Inches(2.5), Inches(5.0), Inches(0.4), "Centroid of Area", size=14, bold=True, color=RGBColor(0x7D,0xD3,0xD8))
    add_textbox(s, Inches(7.35), Inches(2.95), Inches(5.0), Inches(1.1),
                "WPI = ∫ μ_agg(y)·y dy  /  ∫ μ_agg(y) dy", size=17, bold=True, color=WHITE)
    add_pill(s, Inches(7.1), Inches(4.7), Inches(5.5), Inches(0.65), "1000-point discretization for numerical integration", TEAL_LT, text_color=NAVY, size=13)
    add_footer(s, 11, TOTAL)

    # ── Slide 12: Dashboard ─────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    add_header(s, "Real-Time Web Dashboard", subtitle="Flask + Socket.IO · http://127.0.0.1:5050", kicker="Implementation")
    panels = [
        "WPI Arc Gauge — animated 0–10 scale with colour zones",
        "Sensor Cards — live pH / Turbidity / TDS / DO with status bars",
        "WPI Trend Chart — 60-second rolling time series",
        "Membership Visualisation — live μ(x) per linguistic term",
        "Active Rules Panel — which of the 45 rules are firing, and how strongly",
        "Manual Inference — slider-driven one-shot exploration of the FIS",
        "Scenario Selector — Normal, Ideal, Acidic, Turbid, Hard Water, Contaminated",
    ]
    items = [(p, 0, False) for p in panels]
    add_bullets(s, Inches(0.65), Inches(1.9), Inches(11.9), Inches(4.8), items, size=17)
    add_footer(s, 12, TOTAL)

    # ── Slide 13: Validation Methodology ─────────────────────────────
    s = prs.slides.add_slide(blank)
    add_header(s, "Validation Methodology", kicker="Results")
    add_bullets(s, Inches(0.65), Inches(1.85), Inches(11.7), Inches(2.2), [
        ("5,000 synthetic samples labeled by a WHO/BIS hard-threshold oracle (config-defined class balance)", 0, False),
        ("Additionally cross-checked on a real-world Kaggle water-potability dataset", 0, False),
        ("Benchmarked against Random Forest, KNN (k=7), and SVM (RBF) classifiers", 0, False),
        ("Metrics: Accuracy, Macro F1-score, per-class precision/recall, confusion matrix", 0, False),
    ], size=16)
    add_pill(s, Inches(0.65), Inches(4.4), Inches(4.6), Inches(0.55), "python main.py validate", NAVY, size=14)
    add_textbox(s, Inches(0.65), Inches(5.2), Inches(11.5), Inches(0.4), "Boundary-Region Advantage", size=16, bold=True, color=TEAL)
    add_bullets(s, Inches(0.65), Inches(5.6), Inches(11.7), Inches(1.4), [
        ("In the MARGINAL zone (WPI 3.5–6.5), hard thresholds produce abrupt, unrealistic jumps — the FIS's graded membership gives smooth, gradual transitions instead.", 0, False),
    ], size=14)
    add_footer(s, 13, TOTAL)

    # ── Slide 14: Benchmark Results (figure) ──────────────────────────
    s = prs.slides.add_slide(blank)
    add_header(s, "Benchmark Results", subtitle="Synthetic (WHO oracle) vs. Kaggle real-world dataset", kicker="Results")
    add_pic_framed(s, f"{FIG_DIR}/benchmark_comparison.png", Inches(1.9), Inches(1.75), width=Inches(9.5))
    add_footer(s, 14, TOTAL)

    # ── Slide 15: Results table + confusion matrix ────────────────────
    s = prs.slides.add_slide(blank)
    add_header(s, "Interpretability vs. Accuracy Trade-off", kicker="Results")
    rows = [
        ["Method", "Accuracy", "F1-Macro", "Interpretable"],
        ["Mamdani FIS (proposed)", "75.3%", "75.3%", "Yes"],
        ["Random Forest", "99.5%", "99.5%", "No"],
        ["KNN (k=7)", "92.3%", "91.0%", "No"],
        ["SVM (RBF)", "95.7%", "95.0%", "No"],
    ]
    add_table(s, Inches(0.65), Inches(1.8), Inches(6.6), Inches(2.7), rows,
              col_widths=[Inches(2.9), Inches(1.3), Inches(1.3), Inches(1.1)], body_size=13, header_size=13)
    add_pic_framed(s, f"{FIG_DIR}/fis_confusion_matrix.png", Inches(7.7), Inches(1.7), width=Inches(4.9))
    add_textbox(s, Inches(0.65), Inches(4.85), Inches(6.6), Inches(1.8),
                "The FIS trades accuracy for full interpretability: every decision traces back "
                "to specific WHO/BIS-grounded rules. For safety-critical and community-use "
                "deployments, this transparency outweighs the raw-accuracy gap versus black-box models.",
                size=14, color=GREY, line_spacing=1.25)
    add_footer(s, 15, TOTAL)

    # ── Slide 16: Conclusion & Future Scope ───────────────────────────
    s = prs.slides.add_slide(blank)
    add_background(s, NAVY)
    add_textbox(s, Inches(0.7), Inches(0.55), Inches(10), Inches(0.8), "Conclusion & Future Scope", size=30, bold=True, color=WHITE)
    rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.3), Inches(2.2), Pt(3))
    rule.fill.solid(); rule.fill.fore_color.rgb = ACCENT; rule.line.fill.background(); rule.shadow.inherit = False
    add_bullets(s, Inches(0.7), Inches(1.65), Inches(11.9), Inches(2.6), [
        ("A Mamdani FIS on 4 low-cost sensors delivers real-time, interpretable water potability classification", 0, True),
        ("Every decision is traceable to WHO/BIS-grounded rules — critical for safety and regulatory trust", 0, True),
        ("Deployable end-to-end on a ~₹5,350 sensor stack + ESP32, with a live web dashboard", 0, True),
    ], size=17, color=WHITE)
    add_textbox(s, Inches(0.7), Inches(4.4), Inches(6), Inches(0.4), "Future Scope", size=17, bold=True, color=ACCENT)
    add_bullets(s, Inches(0.7), Inches(4.85), Inches(11.9), Inches(2.0), [
        ("Adaptive/learned membership functions via hybrid ANFIS tuning", 0, False),
        ("Field deployment & calibration across multiple real water sources", 0, False),
        ("Edge-native inference directly on ESP32 (no PC bridge required)", 0, False),
    ], size=15, color=RGBColor(0xD6, 0xE6, 0xE8))
    add_textbox(s, Inches(0.7), Inches(6.9), Inches(10), Inches(0.4),
                "Thank you  ·  Standards: WHO Guidelines 2022, BIS IS 10500:2012", size=13, italic=True, color=RGBColor(0x9F,0xC9,0xCC))

    os.makedirs("outputs", exist_ok=True)
    out_path = "outputs/Water_Potability_FIS_Presentation.pptx"
    prs.save(out_path)
    print(f"[*] Presentation saved to {out_path}  ({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    create_ppt()
